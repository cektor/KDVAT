from PyQt5.QtWidgets import *
from PyQt5.QtCore import *
from PyQt5.QtGui import *
from PyQt5.QtPrintSupport import QPrinter, QPrintDialog
import sys
import json
import datetime
from PyQt5.QtCore import QSettings
from openpyxl import Workbook
import requests
import matplotlib.pyplot as plt
from matplotlib.backends.backend_qt5agg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.figure import Figure
from datetime import datetime, timedelta
import os
from PyQt5.QtGui import QIcon
import platform

def get_app_data_path():
    """Platform bağımsız uygulama veri dizini yolu"""
    if platform.system() == "Windows":
        return os.path.join(os.getenv('APPDATA'), 'KDVAT')
    elif platform.system() == "Darwin":  # macOS
        return os.path.expanduser('~/Library/Application Support/KDVAT')
    else:  # Linux ve diğerleri
        return os.path.expanduser('~/.local/share/KDVAT')

def ensure_app_dirs():
    """Uygulama dizinlerinin varlığını kontrol et ve oluştur"""
    app_dir = get_app_data_path()
    if not os.path.exists(app_dir):
        os.makedirs(app_dir)
    return app_dir

def get_resource_path(filename):
    """Kaynak dosyaların platformdan bağımsız yolunu al"""
    if hasattr(sys, '_MEIPASS'):  # PyInstaller ile paketlenmiş
        return os.path.join(sys._MEIPASS, filename)
    
    paths = [
        os.path.join(os.path.dirname(__file__), filename),  # Geliştirme ortamı
        os.path.join(os.getcwd(), filename),  # Çalışma dizini
        os.path.join('/usr/share/kdvat', filename),  # Linux sistem
        os.path.join('/usr/local/share/kdvat', filename),  # macOS sistem
        os.path.join(os.getenv('PROGRAMFILES', ''), 'KDVAT', filename),  # Windows
    ]
    
    for path in paths:
        if os.path.exists(path):
            return path
    return None

class ExchangeRates:
    def __init__(self):
        self.base_url = "https://api.exchangerate-api.com/v4/latest/TRY"
        self.rates = {
            'USD': 0.032,
            'EUR': 0.029,
            'GBP': 0.025,
            'TRY': 1.0,
            'JPY': 4.76,
            'CHF': 0.028,
            'AUD': 0.049,
            'CAD': 0.043
        }
        self.last_update = None
        self.update_rates()
    
    def update_rates(self):
        try:
            response = requests.get(self.base_url)
            data = response.json()
            self.rates = data['rates']
            self.last_update = datetime.now()
        except:
            # API hatası durumunda varsayılan kurları kullan
            pass
    
    def convert(self, amount, from_currency, to_currency):
        try:
            # Kurlar 1 saatten eskiyse güncelle
            if self.last_update and (datetime.now() - self.last_update).total_seconds() > 3600:
                self.update_rates()
                
            if from_currency == to_currency:
                return amount
                
            rate = self.rates[to_currency] / self.rates[from_currency]
            return amount * rate
        except:
            return amount

class KDVHesaplama(QMainWindow):
    def __init__(self):
        super().__init__()
        ensure_app_dirs()  # Uygulama dizinlerini kontrol et
        self.exchange = ExchangeRates()
        self.settings = QSettings('ALGYazilim', 'KDVAT')
        
        # Platform özel ayarlar
        if platform.system() == "Darwin":  # macOS
            self.setUnifiedTitleAndToolBarOnMac(True)
        
        # Çeviri sözlüklerini tanımla
        self.translations = {
            'tr': {
                'window_title': 'KDVAT | KDV Hesaplama Aracı',
                'menu_file': 'Dosya',
                'menu_save': 'Kaydet',
                'menu_load': 'Yükle',
                'menu_print': 'Yazdır',
                'menu_exit': 'Çıkış',
                'menu_settings': 'Ayarlar',
                'menu_vat_management': 'KDV Oran Yönetimi',
                'menu_tools': 'Araçlar',
                'menu_export_excel': "Excel'e Aktar",
                'menu_statistics': 'İstatistikler',
                'menu_bulk_calc': 'Toplu Hesaplama',
                'menu_view': 'Görünüm',
                'menu_theme': 'Tema',
                'menu_theme_dark': 'Karanlık',
                'menu_theme_light': 'Aydınlık',
                'menu_language': 'Dil',
                'menu_language_tr': 'Türkçe',
                'menu_language_en': 'İngilizce',
                'menu_help': 'Yardım',
                'menu_about': 'Hakkında',
                'amount': 'İşlem Tutarı:',
                'amount_placeholder': 'İşlem Tutarını Giriniz',
                'vat_rate': 'KDV Oranı:',
                'currency': 'Para Birimi:',
                'calculate': 'Hesapla',
                'vat_included': 'KDV Dahil',
                'vat_excluded': 'KDV Hariç',
                'transaction_amount': 'İşlem Tutarı:',
                'vat_amount': 'KDV Tutarı:',
                'total_amount': 'Toplam Tutar:',
                'history': 'İşlem Geçmişi',
                'clear_history': 'Geçmişi Temizle',
                'date': 'Tarih',
                'bulk_calc_title': 'Toplu Hesaplama',
                'bulk_calc_placeholder': 'Her satıra bir tutar giriniz...',
                'totals': 'Toplamlar',
                'total_vat': 'Toplam KDV:',
                'grand_total': 'Genel Toplam:',
                'export_pdf': "PDF'e Aktar",
                'success': 'Başarılı',
                'pdf_saved': 'PDF dosyası kaydedildi!',
                'error': 'Hata',
                'enter_valid_number': 'Geçerli bir sayı giriniz!',
                'ready': 'Hazır',
                'data_saved': 'Veriler kaydedildi',
                'data_loaded': 'Veriler yüklendi',
                'history_cleared': 'Geçmiş temizlendi',
                'save': 'Kaydet',
                'load': 'Yükle',
                'print': 'Yazdır',
                'vat_management': 'KDV Oran Yönetimi',
                'vat_management_title': 'KDV Oran Yönetimi',
                'current_rates': 'Mevcut KDV Oranları:',
                'new_rate': 'Yeni KDV Oranı',
                'add': 'Ekle',
                'delete': 'Seçili Oranı Sil',
                'statistics': 'İstatistikler',
                'total_transactions': 'Toplam İşlem Sayısı:',
                'total_vat_amount': 'Toplam KDV Tutarı:',
                'highest_transaction': 'En Yüksek İşlem:',
                'view': 'Görünüm',
                'theme': 'Tema',
                'dark': 'Karanlık',
                'light': 'Aydınlık',
                'about': 'Hakkında',
                'about_title': 'KDVAT Hakkında',
                'about_desc': 'KDV Hesaplama Aracı',
                'features': 'Özellikler:',
                'developer': 'Geliştirici:',
                'version': 'Sürüm:',
                'date': 'Tarih',
                'amount': 'İşlem Tutarı',
                'currency': 'Para Birimi',
                'vat_rate': 'KDV Oranı',
                'vat_amount': 'KDV Tutarı',
                'total': 'Toplam'
            },
            'en': {
                'window_title': 'KDVAT | VAT Tax Calculator',
                'menu_file': 'File',
                'menu_save': 'Save',
                'menu_load': 'Load',
                'menu_print': 'Print',
                'menu_exit': 'Exit',
                'menu_settings': 'Settings',
                'menu_vat_management': 'VAT Rate Management',
                'menu_tools': 'Tools',
                'menu_export_excel': 'Export to Excel',
                'menu_statistics': 'Statistics',
                'menu_bulk_calc': 'Bulk Calculate',
                'menu_view': 'View',
                'menu_theme': 'Theme',
                'menu_theme_dark': 'Dark',
                'menu_theme_light': 'Light',
                'menu_language': 'Language',
                'menu_language_tr': 'Turkish',
                'menu_language_en': 'English',
                'menu_help': 'Help',
                'menu_about': 'About',
                'amount': 'Amount:',
                'amount_placeholder': 'Enter amount',
                'vat_rate': 'VAT Rate:',
                'currency': 'Currency:',
                'calculate': 'Calculate',
                'vat_included': 'VAT Included',
                'vat_excluded': 'VAT Excluded',
                'transaction_amount': 'Transaction Amount:',
                'vat_amount': 'VAT Amount:',
                'total_amount': 'Total Amount:',
                'history': 'Transaction History',
                'clear_history': 'Clear History',
                'date': 'Date',
                'bulk_calc_title': 'Bulk Calculate',
                'bulk_calc_placeholder': 'Enter one amount per line...',
                'totals': 'Totals',
                'total_vat': 'Total VAT:',
                'grand_total': 'Grand Total:',
                'export_pdf': 'Export to PDF',
                'success': 'Success',
                'pdf_saved': 'PDF file saved!',
                'error': 'Error',
                'enter_valid_number': 'Please enter a valid number!',
                'ready': 'Ready',
                'data_saved': 'Data saved',
                'data_loaded': 'Data loaded',
                'history_cleared': 'History cleared',
                'save': 'Save',
                'load': 'Load',
                'print': 'Print',
                'vat_management': 'VAT Rate Management',
                'vat_management_title': 'VAT Rate Management',
                'current_rates': 'Current VAT Rates:',
                'new_rate': 'New VAT Rate',
                'add': 'Add',
                'delete': 'Delete Selected Rate',
                'statistics': 'Statistics',
                'total_transactions': 'Total Transactions:',
                'total_vat_amount': 'Total VAT Amount:',
                'highest_transaction': 'Highest Transaction:',
                'view': 'View',
                'theme': 'Theme',
                'dark': 'Dark',
                'light': 'Light',
                'about': 'About',
                'about_title': 'About VATT',
                'about_desc': 'VAT Tax Calculator',
                'features': 'Features:',
                'developer': 'Developer:',
                'version': 'Version:',
                'date': 'Date',
                'amount': 'Amount',
                'currency': 'Currency',
                'vat_rate': 'VAT Rate',
                'vat_amount': 'VAT Amount',
                'total': 'Total'
            }
        }
        
        # Dil ayarını yükle
        self.current_language = self.settings.value('language', 'tr')
        
        # UI'ı oluştur ve metinleri güncelle
        self.setup_ui()
        self.update_texts()
        self.exchange = ExchangeRates()
        
    def setup_ui(self):
        self.settings = QSettings('ALGYazilim', 'KDVAT')
        self.setWindowTitle("KDVAT | KDV Hesaplama Aracı")
        
        # Logo ve ikon ayarla
        icon_path = self.get_icon_path()
        if (icon_path):
            self.setWindowIcon(QIcon(icon_path))
        
        # Sabit form boyutu ayarla
        self.setFixedSize(1000, 600)
        
        # Pencere özelliklerini ayarla
        self.setWindowFlags(Qt.Window | Qt.MSWindowsFixedSizeDialogHint)
        
        # Menü oluştur
        self.create_menu()
        
        # Toolbar oluştur
        self.create_toolbar()
        
        # Durum çubuğu
        self.statusBar = self.statusBar()
        self.statusBar.showMessage('Hazır')
        
        # Ana widget
        self.central_widget = QWidget()
        self.setCentralWidget(self.central_widget)
        self.layout = QHBoxLayout(self.central_widget)
        
        # Sol Panel (Hesaplama)
        self.left_panel = QVBoxLayout()
        
        
        
        # Tutar girişi
        self.tutar_label = QLabel("İşlem Tutarı:")
        self.tutar_input = QLineEdit()
        self.tutar_input.setPlaceholderText("İşlem Tutarını giriniz")
        
        # KDV oranı
        self.kdv_label = QLabel("KDV Oranı:")
        self.kdv_combo = QComboBox()
        self.load_kdv_rates()
        
        # Para birimi seçimi
        currency_layout = QHBoxLayout()
        currency_layout.addWidget(QLabel("Para Birimi:"))
        self.currency_combo = QComboBox()
        self.currency_combo.addItems(['TRY', 'USD', 'EUR', 'GBP', 'JPY', 'CHF', 'AUD', 'CAD'])
        
        # Son seçilen para birimini yükle
        last_currency = self.settings.value('last_currency', 'TRY')
        index = self.currency_combo.findText(last_currency)
        if index >= 0:
            self.currency_combo.setCurrentIndex(index)
        
        currency_layout.addWidget(self.currency_combo)
        self.left_panel.addLayout(currency_layout)
        
        # Para birimi değişikliğini kaydet
        self.currency_combo.currentTextChanged.connect(self.save_last_currency)
        
        # Hesapla butonu
        self.hesapla_btn = QPushButton("Hesapla")
        self.hesapla_btn.setShortcut("Ctrl+Return")
        
        # KDV Dahil Sonuçlar
        self.kdv_dahil_group = QGroupBox("KDV Dahil")
        kdv_dahil_layout = QVBoxLayout()
        
        self.islem_tutari_dahil = QLabel("İşlem Tutarı: ")
        self.kdv_tutari_dahil = QLabel("KDV Tutarı: ")
        self.toplam_tutar_dahil = QLabel("Toplam Tutar: ")
        
        kdv_dahil_layout.addWidget(self.islem_tutari_dahil)
        kdv_dahil_layout.addWidget(self.kdv_tutari_dahil)
        kdv_dahil_layout.addWidget(self.toplam_tutar_dahil)
        self.kdv_dahil_group.setLayout(kdv_dahil_layout)
        
        # KDV Hariç Sonuçlar
        self.kdv_haric_group = QGroupBox("KDV Hariç")
        kdv_haric_layout = QVBoxLayout()
        
        self.islem_tutari_haric = QLabel("İşlem Tutarı: ")
        self.kdv_tutari_haric = QLabel("KDV Tutarı: ")
        self.toplam_tutar_haric = QLabel("Toplam Tutar: ")
        
        kdv_haric_layout.addWidget(self.islem_tutari_haric)
        kdv_haric_layout.addWidget(self.kdv_tutari_haric)
        kdv_haric_layout.addWidget(self.toplam_tutar_haric)
        self.kdv_haric_group.setLayout(kdv_haric_layout)
        
        # Layout'a widget'ları ekle
    
        self.left_panel.addWidget(self.tutar_label)
        self.left_panel.addWidget(self.tutar_input)
        self.left_panel.addWidget(self.kdv_label)
        self.left_panel.addWidget(self.kdv_combo)
        self.left_panel.addWidget(self.hesapla_btn)
        self.left_panel.addWidget(self.kdv_dahil_group)
        self.left_panel.addWidget(self.kdv_haric_group)
        
        # Ana layout'a sol paneli ekle
        left_widget = QWidget()
        left_widget.setLayout(self.left_panel)
        self.layout.addWidget(left_widget)
        
        # Sağ Panel (Geçmiş)
        self.right_panel = QVBoxLayout()
        self.history_label = QLabel("İşlem Geçmişi")
        self.history_table = QTableWidget()
        self.history_table.setColumnCount(6)  # Sütun sayısını 6'ya çıkar
        self.history_table.setHorizontalHeaderLabels(
            ["Tarih", "Tutar", "Para Birimi", "KDV Oranı", "KDV Tutarı", "Toplam"])
        
        # Geçmiş işlemleri temizle
        self.clear_history_btn = QPushButton("Geçmişi Temizle")
        
        # Layout'ları düzenle
        self.setup_layout()
        
        # Sinyal bağlantıları
        self.connect_signals()
        
        # Geçmiş işlemleri yükle
        self.load_history()
        
        self.kdv_combo.currentTextChanged.connect(self.save_last_kdv_rate)
        self.load_history_from_settings()

        # Tema stillerini tanımla
        self.dark_theme_style = """
            QMainWindow, QWidget {
                background-color: #1e1e1e;
                color: #ffffff;
            }
            QLabel {
                color: #ffffff;
                font-size: 12px;
            }
            QLineEdit, QComboBox {
                padding: 8px;
                border: 1px solid #3d3d3d;
                border-radius: 4px;
                background-color: #2d2d2d;
                color: #ffffff;
            }
            QPushButton {
                background-color: #0078d4;
                color: white;
                padding: 8px 16px;
                border-radius: 4px;
                border: none;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #1084d8;
            }
            QGroupBox {
                border: 1px solid #3d3d3d;
                border-radius: 6px;
                margin-top: 12px;
                color: #ffffff;
            }
            QTableWidget {
                background-color: #2d2d2d;
                color: #ffffff;
                gridline-color: #3d3d3d;
            }
        """
        
        self.light_theme_style = """
            QMainWindow, QWidget {
                background-color: #f0f0f0;
                color: #000000;
            }
            QLabel {
                color: #000000;
                font-size: 12px;
            }
            QLineEdit, QComboBox {
                padding: 8px;
                border: 1px solid #cccccc;
                border-radius: 4px;
                background-color: #ffffff;
                color: #000000;
            }
            QPushButton {
                background-color: #0078D4;
                color: white;
                padding: 8px 16px;
                border-radius: 4px;
            }
            QPushButton:hover {
                background-color: #1084d8;
            }
            QGroupBox {
                border: 1px solid #cccccc;
                border-radius: 6px;
                margin-top: 12px;
                color: #000000;
            }
            QTableWidget {
                background-color: #ffffff;
                color: #000000;
                gridline-color: #cccccc;
            }
        """
        
        # Varsayılan tema ayarla
        self.setStyleSheet(self.dark_theme_style)

        # Son kullanılan temayı yükle
        last_theme = self.settings.value('theme', 'dark')
        self.change_theme(last_theme)

        # Dil menüsü ekle
        language_menu = self.menuBar().addMenu(self.tr('Dil'))
        tr_action = QAction('Türkçe', self)
        en_action = QAction('English', self)
        tr_action.triggered.connect(lambda: self.change_language('tr'))
        en_action.triggered.connect(lambda: self.change_language('en'))
        language_menu.addAction(tr_action)
        language_menu.addAction(en_action)

        # İşlem geçmişi tablosu ayarları
        self.history_table.setEditTriggers(QTableWidget.NoEditTriggers)  # ReadOnly

        # Seçili satırı silme butonu
        self.delete_row_btn = QPushButton("Seçili Satırı Sil")
        self.delete_row_btn.setStyleSheet("""
            QPushButton {
                background-color: #f44336;
                color: white;
                padding: 8px 16px;
                border-radius: 4px;
                border: none;
            }
            QPushButton:hover {
                background-color: #d32f2f;
            }
        """)
        self.delete_row_btn.clicked.connect(self.delete_selected_row)

        # Layout güncellemeleri
        self.right_panel.insertWidget(self.right_panel.count()-1, self.delete_row_btn)

        # Geçmişi temizle butonu stil ekle
        self.clear_history_btn.setStyleSheet("""
            QPushButton {
                background-color: #f44336;
                color: white;
                padding: 8px 16px;
                border-radius: 4px;
                border: none;
            }
            QPushButton:hover {
                background-color: #d32f2f;
            }
        """)

    def create_menu(self):
        menubar = self.menuBar()
        
        # Dosya menüsü
        file_menu = menubar.addMenu('Dosya')
        
        save_action = QAction('Kaydet', self)
        save_action.setShortcut('Ctrl+S')
        save_action.triggered.connect(self.save_data)
        
        load_action = QAction('Yükle', self)
        load_action.setShortcut('Ctrl+O')
        load_action.triggered.connect(self.load_data)
        
        print_action = QAction('Yazdır', self)
        print_action.setShortcut('Ctrl+P')
        print_action.triggered.connect(self.print_data)
        
        exit_action = QAction('Çıkış', self)
        exit_action.setShortcut('Ctrl+Q')
        exit_action.triggered.connect(self.close)
        
        file_menu.addAction(save_action)
        file_menu.addAction(load_action)
        file_menu.addAction(print_action)
        file_menu.addSeparator()
        file_menu.addAction(exit_action)
        
        # Ayarlar menüsü
        settings_menu = menubar.addMenu('Ayarlar')
        
        kdv_yonetim_action = QAction('KDV Oran Yönetimi', self)
        kdv_yonetim_action.setShortcut('Ctrl+K')
        kdv_yonetim_action.triggered.connect(self.show_kdv_yonetim)
        
        settings_menu.addAction(kdv_yonetim_action)

        # Araçlar menüsü
        tools_menu = menubar.addMenu('Araçlar')
        
        excel_export = QAction('Excel\'e Aktar', self)
        excel_export.setShortcut('Ctrl+E')
        excel_export.triggered.connect(self.export_to_excel)
        
        stats_action = QAction('İstatistikler', self)
        stats_action.setShortcut('Ctrl+I')
        stats_action.triggered.connect(self.show_statistics)
        
        bulk_calc = QAction('Toplu Hesaplama', self)
        bulk_calc.setShortcut('Ctrl+B')
        bulk_calc.triggered.connect(self.show_bulk_calc)
        
        tools_menu.addAction(excel_export)
        tools_menu.addAction(stats_action)
        tools_menu.addAction(bulk_calc)
        
        # Yardım menüsü
        help_menu = menubar.addMenu('Yardım')
        about_action = QAction('Hakkında', self)
        about_action.triggered.connect(self.show_about)
        help_menu.addAction(about_action)

        # Dil menüsü
        lang_menu = menubar.addMenu(self.tr('menu_language'))
        
        tr_action = QAction('Türkçe', self)
        tr_action.setIcon(QIcon('/usr/share/icons/hicolor/48x48/apps/tr.png'))
        tr_action.triggered.connect(lambda: self.change_language('tr'))
        
        en_action = QAction('English', self)
        en_action.setIcon(QIcon('/usr/share/icons/hicolor/48x48/apps/en.png'))
        en_action.triggered.connect(lambda: self.change_language('en'))
        
        lang_menu.addAction(tr_action)
        lang_menu.addAction(en_action)
        
        # Son seçilen dili işaretle
        current_lang = self.settings.value('language', 'tr')
        if current_lang == 'tr':
            tr_action.setChecked(True)
        else:
            en_action.setChecked(True)

    def show_about(self):
        about_dialog = QDialog(self)
        about_dialog.setWindowTitle("Hakkında")
        about_dialog.setFixedSize(500, 630)
        
        layout = QVBoxLayout()
        
        # Logo ekle
        logo_path = self.get_logo_path()
        if logo_path:
            logo_label = QLabel()
            pixmap = QPixmap(logo_path)
            scaled_pixmap = pixmap.scaled(100, 100, Qt.KeepAspectRatio, Qt.SmoothTransformation)
            logo_label.setPixmap(scaled_pixmap)
            logo_label.setAlignment(Qt.AlignCenter)
            layout.addWidget(logo_label)
        
        # Başlık
        title_label = QLabel("KDVAT")
        title_label.setStyleSheet("font-size: 15px; font-weight: bold; color: #2196F3;")
        title_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(title_label)
        
        # Bilgi metni
        info_text = """
        <p>KDVAT | KDV Hesaplama Aracı</p>
        <p>Bu program, KDV hesaplamalarını kolaylaştırmak için geliştirilmiştir.</p>
        <br>
        <p><b>Özellikler:</b></p>
        <ul>
            <li>KDV Dahil/Hariç hesaplama</li>
            <li>Toplu hesaplama</li>
            <li>Excel'e aktarma</li>
            <li>İstatistik raporları</li>
            <li>İşlem geçmişi</li>

        </ul>
        <br>
        <p>Geliştirici: ALG Yazılım Inc.©</p>
        <p>www.algyazilim.com | info@algyazilim.com</p>
        <p>Fatih ÖNDER (CekToR) | fatih@algyazilim.com</p>
        <p>GitHub: https://github.com/cektor</p>
        <p>Sürüm: 1.0</p>
        <p>ALG Yazılım Pardus'a Göç'ü Destekler.</p>
        <p>Telif Hakkı © 2025 GNU</p>

        """
        
        info_label = QLabel(info_text)
        info_label.setWordWrap(True)
        info_label.setOpenExternalLinks(True)
        info_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(info_label)
        
        about_dialog.setLayout(layout)
        
        # Dialog'u tema ile uyumlu hale getir
        if self.settings.value('theme', 'dark') == 'dark':
            about_dialog.setStyleSheet("""
                QDialog {
                    background-color: #1e1e1e;
                    color: #ffffff;
                }
                QLabel {
                    color: #ffffff;
                }
            """)
        else:
            about_dialog.setStyleSheet("""
                QDialog {
                    background-color: #f0f0f0;
                    color: #000000;
                }
                QLabel {
                    color: #000000;
                }
            """)
        
        about_dialog.exec_()

    def show_kdv_yonetim(self):
        dialog = KDVOranYonetimi(self)
        dialog.exec_()

    def create_toolbar(self):
        toolbar = QToolBar()
        self.addToolBar(toolbar)
        
        save_action = QAction('Kaydet', self)
        save_action.triggered.connect(self.save_data)
        toolbar.addAction(save_action)
        
        load_action = QAction('Yükle', self)
        load_action.triggered.connect(self.load_data)
        toolbar.addAction(load_action)
        
        print_action = QAction('Yazdır', self)
        print_action.triggered.connect(self.print_data)
        toolbar.addAction(print_action)

    def save_data(self):
        filename, _ = QFileDialog.getSaveFileName(self, "Kaydet", "", "JSON Dosyası (*.json)")
        if filename:
            # .json uzantısını kontrol et ve ekle
            if not filename.endswith('.json'):
                filename += '.json'
                
            data = {
                'history': self.get_history_data()
            }
            with open(filename, 'w') as f:
                json.dump(data, f)
            self.statusBar.showMessage('Veriler kaydedildi')

    def load_data(self):
        filename, _ = QFileDialog.getOpenFileName(self, "Yükle", "", "JSON Dosyası (*.json)")
        if filename:
            try:
                with open(filename, 'r') as f:
                    data = json.load(f)
                    self.load_history_data(data['history'])
                    self.update_totals()  # Toplamları güncelle
                self.statusBar.showMessage('Veriler yüklendi')
            except Exception as e:
                QMessageBox.warning(self, "Hata", f"Dosya yüklenirken hata: {str(e)}")

    def print_data(self):
        printer = QPrinter(QPrinter.HighResolution)
        dialog = QPrintDialog(printer, self)
        
        # Dialog başlığını ayarla
        dialog.setWindowTitle("Yazdır")
        
        # Butonları Türkçeleştir
        for button in dialog.findChildren(QPushButton):
            if button.text() == "&Print":
                button.setText("&Yazdır")
            elif button.text() == "&Cancel":
                button.setText("&İptal")
            elif button.text() == "Properties":
                button.setText("Özellikler")
                
        # Yazdırma seçeneklerini Türkçeleştir
        for child in dialog.findChildren(QWidget):
            if hasattr(child, 'setText'):
                if child.text() == "Print range":
                    child.setText("Yazdırma aralığı")
                elif child.text() == "All":
                    child.setText("Tümü")
                elif child.text() == "Selection":
                    child.setText("Seçili")
                elif child.text() == "Pages":
                    child.setText("Sayfalar")
                elif child.text() == "Copies":
                    child.setText("Kopya sayısı")
        
        if dialog.exec_() == QPrintDialog.Accepted:
            self.print_document(printer)

    def add_to_history(self, tutar, kdv_orani, kdv_tutari, toplam):
        currency = self.currency_combo.currentText()
        row = self.history_table.rowCount()
        self.history_table.insertRow(row)
        
        self.history_table.setItem(row, 0, QTableWidgetItem(datetime.now().strftime("%Y-%m-%d %H:%M")))
        self.history_table.setItem(row, 1, QTableWidgetItem(f"{self.format_currency(tutar, currency)}"))
        self.history_table.setItem(row, 2, QTableWidgetItem(currency))
        self.history_table.setItem(row, 3, QTableWidgetItem(f"%{kdv_orani*100}"))
        self.history_table.setItem(row, 4, QTableWidgetItem(f"{self.format_currency(kdv_tutari, currency)}"))
        self.history_table.setItem(row, 5, QTableWidgetItem(f"{self.format_currency(toplam, currency)}"))
        self.update_totals()

    def hesapla(self):
        try:
            tutar = float(self.tutar_input.text() or 0)
            kdv_orani = float(self.kdv_combo.currentText().strip('%')) / 100
            currency = self.currency_combo.currentText()
            
            # KDV Dahil Hesaplama
            kdv_tutari_dahil = tutar * kdv_orani
            toplam_tutar_dahil = tutar + kdv_tutari_dahil
            
            self.islem_tutari_dahil.setText(f"İşlem Tutarı: {self.format_currency(tutar, currency)} {currency}")
            self.kdv_tutari_dahil.setText(f"KDV Tutarı: {self.format_currency(kdv_tutari_dahil, currency)} {currency}")
            self.toplam_tutar_dahil.setText(f"Toplam Tutar: {self.format_currency(toplam_tutar_dahil, currency)} {currency}")
            
            # KDV Hariç Hesaplama
            islem_tutari_haric = tutar / (1 + kdv_orani)
            kdv_tutari_haric = tutar - islem_tutari_haric
            
            self.islem_tutari_haric.setText(f"İşlem Tutarı: {self.format_currency(islem_tutari_haric, currency)} {currency}")
            self.kdv_tutari_haric.setText(f"KDV Tutarı: {self.format_currency(kdv_tutari_haric, currency)} {currency}")
            self.toplam_tutar_haric.setText(f"Toplam Tutar: {self.format_currency(tutar, currency)} {currency}")
            
            # İşlem geçmişine ekle - formatlanmış değerler
            self.add_to_history(tutar, kdv_orani, kdv_tutari_dahil, toplam_tutar_dahil)
            
            self.tutar_input.clear()
                
        except ValueError as e:
            QMessageBox.warning(self, "Hata", str(e))
        except Exception as e:
            QMessageBox.critical(self, "Hata", f"Beklenmeyen bir hata oluştu: {str(e)}")

    def kdv_oran_ekle(self):
        try:
            yeni_oran = float(self.yeni_oran_input.text())
            if 0 < yeni_oran <= 100:
                yeni_oran_text = f"%{yeni_oran}"
                if yeni_oran_text not in [self.kdv_combo.itemText(i) for i in range(self.kdv_combo.count())]:
                    self.kdv_combo.addItem(yeni_oran_text)
                    self.yeni_oran_input.clear()
                    self.save_kdv_rates()  # Yeni oranı kaydet
                    self.statusBar.showMessage(f'Yeni KDV oranı eklendi: {yeni_oran_text}')
        except ValueError:
            QMessageBox.warning(self, "Hata", "Geçerli bir sayı giriniz!")

    def kdv_oran_sil(self):
        current_index = self.kdv_combo.currentIndex()
        if current_index >= 0:
            removed_item = self.kdv_combo.currentText()
            self.kdv_combo.removeItem(current_index)
            self.save_kdv_rates()  # Güncel oranları kaydet
            self.statusBar.showMessage(f'KDV oranı silindi: {removed_item}')

    def get_history_data(self):
        history = []
        for row in range(self.history_table.rowCount()):
            row_data = {
                'tarih': self.history_table.item(row, 0).text(),
                'tutar': float(self.history_table.item(row, 1).text()),
                'para_birimi': self.history_table.item(row, 2).text(),
                'kdv_orani': self.history_table.item(row, 3).text(),
                'kdv_tutari': float(self.history_table.item(row, 4).text()),
                'toplam': float(self.history_table.item(row, 5).text())
            }
            history.append(row_data)
        return history

    def load_history_data(self, history):
        self.history_table.setRowCount(0)
        if isinstance(history, list):
            for item in history:
                row = self.history_table.rowCount()
                self.history_table.insertRow(row)
                self.history_table.setItem(row, 0, QTableWidgetItem(item.get('tarih', '')))
                self.history_table.setItem(row, 1, QTableWidgetItem(f"{item.get('tutar', ''):.2f}"))
                self.history_table.setItem(row, 2, QTableWidgetItem(item.get('para_birimi', '')))
                self.history_table.setItem(row, 3, QTableWidgetItem(item.get('kdv_orani', '')))
                self.history_table.setItem(row, 4, QTableWidgetItem(f"{item.get('kdv_tutari', ''):.2f}"))
                self.history_table.setItem(row, 5, QTableWidgetItem(f"{item.get('toplam', ''):.2f}"))
            
            # Verileri yükledikten sonra toplamları güncelle
            self.update_totals()

    def print_document(self, printer):
        document = QTextDocument()
        cursor = QTextCursor(document)
        
        # Başlık
        format = cursor.charFormat()
        format.setFontPointSize(14)
        cursor.insertText("KDV Hesaplama Raporu\n\n", format)
        
        # Tablo başlıkları
        headers = ["Tarih", "Tutar", "Para Birimi", "KDV Oranı", "KDV Tutarı", "Toplam"]
        table = cursor.insertTable(self.history_table.rowCount() + 1, len(headers))
        
        for col, header in enumerate(headers):
            cursor.insertText(header)
            cursor.movePosition(QTextCursor.NextCell)
        
        # Tablo verileri
        for row in range(self.history_table.rowCount()):
            for col in range(len(headers)):
                cursor.insertText(self.history_table.item(row, col).text())
                cursor.movePosition(QTextCursor.NextCell)
        
        # Toplamları ekle
        cursor.movePosition(QTextCursor.End)
        cursor.insertBlock()
        cursor.insertBlock()
        
        toplam_tutar, toplam_kdv, genel_toplam = self.calculate_totals()
        currency = self.currency_combo.currentText()
        
        format.setFontPointSize(12)
        cursor.setCharFormat(format)
        
        cursor.insertText(f"Toplam KDV Hariç Tutar: {self.format_currency(toplam_tutar, currency)} {currency}\n")
        cursor.insertText(f"Toplam KDV Tutarı: {self.format_currency(toplam_kdv, currency)} {currency}\n")
        cursor.insertText(f"Genel Toplam KDV Dahil: {self.format_currency(genel_toplam, currency)} {currency}")
        
        document.print_(printer)

    def setup_layout(self):
        # Sol panel bileşenleri

      
        self.left_panel.addWidget(self.tutar_label)
        self.left_panel.addWidget(self.tutar_input)
        self.left_panel.addWidget(self.kdv_label)
        self.left_panel.addWidget(self.kdv_combo)
        self.left_panel.addWidget(self.hesapla_btn)
        self.left_panel.addStretch()
        
        # Layout'a sonuç gruplarını ekle
        self.left_panel.addWidget(self.kdv_dahil_group)
        self.left_panel.addWidget(self.kdv_haric_group)
        
        # Sağ panel bileşenleri
        self.right_panel.addWidget(self.history_label)
        self.right_panel.addWidget(self.history_table)
        
        # Toplamlar için grup kutusu
        totals_group = QGroupBox("Toplamlar")
        totals_layout = QVBoxLayout()
        
        self.total_amount_label = QLabel("Toplam KDV Hariç Tutar: 0.00")
        self.total_vat_label = QLabel("Toplam KDV Tutarı: 0.00")
        self.total_sum_label = QLabel("Genel Toplam KDV Dahil: 0.00")
        
        totals_layout.addWidget(self.total_amount_label)
        totals_layout.addWidget(self.total_vat_label)
        totals_layout.addWidget(self.total_sum_label)
        
        totals_group.setLayout(totals_layout)
        self.right_panel.addWidget(totals_group)
        
        # Geçmişi temizle butonu
        self.right_panel.addWidget(self.clear_history_btn)
        
        # Ana layout
        left_widget = QWidget()
        left_widget.setLayout(self.left_panel)
        
        right_widget = QWidget()
        right_widget.setLayout(self.right_panel)
        
        splitter = QSplitter(Qt.Horizontal)
        splitter.addWidget(left_widget)
        splitter.addWidget(right_widget)
        
        self.layout.addWidget(splitter)

    def update_totals(self):
        try:
            toplam_tutar, toplam_kdv, genel_toplam = self.calculate_totals()
            currency = self.currency_combo.currentText()
            
            self.total_amount_label.setText(f"Toplam KDV Hariç Tutar: {self.format_currency(toplam_tutar, currency)} {currency}")
            self.total_vat_label.setText(f"Toplam KDV Tutarı: {self.format_currency(toplam_kdv, currency)} {currency}")
            self.total_sum_label.setText(f"Genel Toplam KDV Dahil: {self.format_currency(genel_toplam, currency)} {currency}")
        except Exception as e:
            print(f"Toplamları güncellerken hata: {str(e)}")

    def connect_signals(self):
        self.hesapla_btn.clicked.connect(self.hesapla)
        self.clear_history_btn.clicked.connect(self.clear_history)
        # Enter tuşu bağlantısı
        self.tutar_input.returnPressed.connect(self.hesapla)
        
        # Otomatik hesaplamayı kaldır
        # self.tutar_input.textChanged.connect(self.hesapla)
        # self.kdv_combo.currentTextChanged.connect(self.hesapla)
 

    def clear_history(self):
        self.history_table.setRowCount(0)
        self.save_history_to_settings()  # Temizlenen geçmişi kaydet
        self.statusBar.showMessage('Geçmiş temizlendi')
        self.update_totals()

    def load_history(self):
        try:
            with open('kdv_history.json', 'r') as f:
                data = json.load(f)
                self.load_history_data(data)
        except FileNotFoundError:
            # İlk çalıştırmada dosya yoksa boş geç
            pass
        except Exception as e:
            self.statusBar.showMessage(f'Geçmiş yüklenirken hata: {str(e)}')

    def load_history_data(self, data):
        if isinstance(data, list):
            for item in data:
                row = self.history_table.rowCount()
                self.history_table.insertRow(row)
                self.history_table.setItem(row, 0, QTableWidgetItem(item.get('tarih', '')))
                self.history_table.setItem(row, 1, QTableWidgetItem(str(item.get('tutar', ''))))
                self.history_table.setItem(row, 2, QTableWidgetItem(item.get('para_birimi', '')))
                self.history_table.setItem(row, 3, QTableWidgetItem(item.get('kdv_orani', '')))
                self.history_table.setItem(row, 4, QTableWidgetItem(str(item.get('kdv_tutari', ''))))
                self.history_table.setItem(row, 5, QTableWidgetItem(str(item.get('toplam', ''))))

    def load_kdv_rates(self):
        saved_rates = self.settings.value('kdv_rates', ['%1', '%8', '%18', '%20'], type=list)
        self.kdv_combo.clear()
        self.kdv_combo.addItems(saved_rates)
        
        # Son seçilen KDV oranını yükle
        last_selected = self.settings.value('last_kdv_rate', '%20', type=str)
        
        if last_selected in saved_rates:
            index = self.kdv_combo.findText(last_selected)
        else:
            index = 0
        
        self.kdv_combo.setCurrentIndex(index)


    def save_kdv_rates(self):
        rates = [self.kdv_combo.itemText(i) for i in range(self.kdv_combo.count())]
        self.settings.setValue('kdv_rates', rates)

    def save_last_kdv_rate(self):
        current_rate = self.kdv_combo.currentText()
        self.settings.setValue('last_kdv_rate', current_rate)
        self.settings.sync()

    def save_last_currency(self):
        current_currency = self.currency_combo.currentText()
        self.settings.setValue('last_currency', current_currency)
        self.settings.sync()

    def closeEvent(self, event):
        # Program kapanırken geçmişi kaydet
        self.save_history_to_settings()
        event.accept()

    def save_history_to_settings(self):
        history_file = os.path.join(get_app_data_path(), 'history.json')
        try:
            history_data = []
            for row in range(self.history_table.rowCount()):
                row_data = {
                    'tarih': self.history_table.item(row, 0).text(),
                    'tutar': self.history_table.item(row, 1).text(),
                    'para_birimi': self.history_table.item(row, 2).text(),
                    'kdv_orani': self.history_table.item(row, 3).text(),
                    'kdv_tutari': self.history_table.item(row, 4).text(),
                    'toplam': self.history_table.item(row, 5).text()
                }
                history_data.append(row_data)
            
            with open(history_file, 'w', encoding='utf-8') as f:
                json.dump(history_data, f, ensure_ascii=False, indent=2)
                
        except Exception as e:
            print(f"Geçmiş kaydedilirken hata: {str(e)}")

    def load_history_from_settings(self):
        history_file = os.path.join(get_app_data_path(), 'history.json')
        try:
            if os.path.exists(history_file):
                with open(history_file, 'r', encoding='utf-8') as f:
                    history_data = json.load(f)
                    
                for item in history_data:
                    row = self.history_table.rowCount()
                    self.history_table.insertRow(row)
                    self.history_table.setItem(row, 0, QTableWidgetItem(item['tarih']))
                    self.history_table.setItem(row, 1, QTableWidgetItem(item['tutar']))
                    self.history_table.setItem(row, 2, QTableWidgetItem(item['para_birimi']))
                    self.history_table.setItem(row, 3, QTableWidgetItem(item['kdv_orani']))
                    self.history_table.setItem(row, 4, QTableWidgetItem(item['kdv_tutari']))
                    self.history_table.setItem(row, 5, QTableWidgetItem(item['toplam']))
                    
        except Exception as e:
            print(f"Geçmiş yüklenirken hata: {str(e)}")
        
        self.update_totals()

    def calculate_totals(self):
        toplam_tutar = 0
        toplam_kdv = 0
        genel_toplam = 0
        
        try:
            for i in range(self.history_table.rowCount()):
                # Doğru sütunlardan değerleri al
                tutar_text = self.history_table.item(i, 1).text().split()[0].replace(',', '.')  # Tutar
                kdv_tutari_text = self.history_table.item(i, 4).text().split()[0].replace(',', '.')  # KDV Tutarı
                toplam_text = self.history_table.item(i, 5).text().split()[0].replace(',', '.')  # Toplam Tutar
                
                # Float dönüşümü
                tutar = float(tutar_text)
                kdv_tutari = float(kdv_tutari_text)
                toplam = float(toplam_text)
                
                toplam_tutar += tutar
                toplam_kdv += kdv_tutari
                genel_toplam += toplam
                
            return toplam_tutar, toplam_kdv, genel_toplam
            
        except (ValueError, AttributeError) as e:
            QMessageBox.warning(self, "Hata", f"Sayısal değer dönüştürme hatası: {str(e)}")
            return 0, 0, 0

    def export_to_excel(self):
        try:
            # Dosya formatları için filtreler
            file_filters = "Excel Çalışma Kitabı (*.xlsx);;Excel Binary Çalışma Kitabı (*.xlsb);;Excel Şablonu (*.xltx);;OpenDocument Hesap Tablosu (*.ods);;Excel Makro Özellikli Şablon (*.xltm);;OpenDocument Hesap Tablosu Şablonu (*.ots);;CSV Dosyası (*.csv);;PDF Dosyası (*.pdf)"
            
            file_name, selected_filter = QFileDialog.getSaveFileName(self, "Excel'e Aktar", "", file_filters)
            
            if file_name:
                if not any(file_name.endswith(ext) for ext in ['.xlsx', '.xlsb', '.xltx', '.ods', '.xltm', '.ots', '.csv', '.pdf']):
                    ext = selected_filter.split('*')[1].split(')')[0]
                    file_name = f"{file_name}{ext}"
                
                if file_name.endswith('.pdf'):
                    self.export_to_pdf(file_name)
                else:
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "KDV Hesaplama"
                    
                    # Sabit sütun başlıkları
                    headers = ['Tarih', 'Tutar', 'Para Birimi', 'KDV Oranı (%)', 'KDV Tutarı', 'Toplam Tutar']
                    ws.append(headers)
                    
                    # Verileri doğru sırayla ekle
                    for i in range(self.history_table.rowCount()):
                        row_data = []
                        for j in range(6):  # 6 sütun
                            item = self.history_table.item(i, j)
                            if item:
                                value = item.text()
                                # KDV oranından % işaretini kaldır
                                if j == 3:  # KDV Oranı sütunu
                                    value = value.replace('%', '').strip()
                                row_data.append(value)
                            else:
                                row_data.append('')
                        ws.append(row_data)
                    
                    # Toplamları hesapla
                    toplam_tutar, toplam_kdv, genel_toplam = self.calculate_totals()
                    
                    # Boş satır ekle
                    ws.append([])
                    
                    # Toplamları ekle
                    ws.append(['Toplam KDV Hariç Tutar:', f'{toplam_tutar:.2f}'])
                    ws.append(['Toplam KDV Tutarı:', f'{toplam_kdv:.2f}'])
                    ws.append(['Genel Toplam KDV Dahil:', f'{genel_toplam:.2f}'])
                    
                    if file_name.endswith('.csv'):
                        with open(file_name, 'w', newline='', encoding='utf-8') as f:
                            for row in ws.rows:
                                f.write(','.join(str(cell.value) for cell in row) + '\n')
                    else:
                        wb.save(file_name)
                    
                QMessageBox.information(self, "Bilgi", f"Veriler başarıyla {file_name} dosyasına kaydedildi.")
        
        except Exception as e:
            QMessageBox.critical(self, "Hata", f"Dosya kaydedilirken hata oluştu: {str(e)}")

    def export_to_pdf(self, file_name):
        printer = QPrinter(QPrinter.HighResolution)
        printer.setOutputFormat(QPrinter.PdfFormat)
        printer.setOutputFileName(file_name)
        
        doc = QTextDocument()
        html = "<h2>KDV Hesaplama Raporu</h2><br><table border='1' cellspacing='0' cellpadding='3'>"
        
        # Sabit sütun başlıkları
        headers = ['Tarih', 'Tutar', 'Para Birimi', 'KDV Oranı (%)', 'KDV Tutarı', 'Toplam Tutar']
        html += "<tr>"
        for header in headers:
            html += f"<th>{header}</th>"
        html += "</tr>"
        
        # Verileri doğru sırayla ekle
        for i in range(self.history_table.rowCount()):
            html += "<tr>"
            for j in range(6):  # 6 sütun
                item = self.history_table.item(i, j)
                value = item.text() if item else ''
                html += f"<td>{value}</td>"
            html += "</tr>"
        
        html += "</table><br><br>"
        
        # Toplamları hesapla ve ekle
        toplam_tutar, toplam_kdv, genel_toplam = self.calculate_totals()
        
        html += f"""
        <p><strong>Toplam KDV Hariç Tutar:</strong> {toplam_tutar:.2f}</p>
        <p><strong>Toplam KDV Tutarı:</strong> {toplam_kdv:.2f}</p>
        <p><strong>Genel Toplam KDV Dahil:</strong> {genel_toplam:.2f}</p>
        """
        
        doc.setHtml(html)
        doc.print_(printer)

    def show_statistics(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("İstatistikler")
        dialog.setModal(True)
        dialog.resize(500, 700)
        
        layout = QVBoxLayout()
        
        # Başlık
        title = QLabel("KDV Hesaplama İstatistikleri")
        title.setStyleSheet("""
            QLabel {
                font-size: 15px;
                font-weight: bold;
                color: #2196F3;
                padding: 10px;
            }
        """)
        title.setAlignment(Qt.AlignCenter)
        layout.addWidget(title)
        
        # İstatistik Kartları
        stats_group = QGroupBox("Genel İstatistikler")
        stats_group.setStyleSheet("""
            QGroupBox {
                background-color: #2D2D2D;
                border: 1px solid #3D3D3D;
                border-radius: 6px;
                margin-top: 10px;
                padding: 10px;
                color: white;
            }
            QGroupBox::title {
                color: #2196F3;
                subcontrol-position: top center;
                padding: 5px;
            }
        """)
        
        stats_layout = QGridLayout()
        
        # İstatistikleri hesapla
        total_transactions = self.history_table.rowCount()
        
        total_amount = sum(float(self.history_table.item(row, 1).text().split()[0].replace(',', '.'))
                          for row in range(total_transactions))
        
        total_vat = sum(float(self.history_table.item(row, 4).text().split()[0].replace(',', '.'))
                        for row in range(total_transactions))
        
        max_amount = max(float(self.history_table.item(row, 5).text().split()[0].replace(',', '.'))
                         for row in range(total_transactions)) if total_transactions > 0 else 0
        
        avg_amount = total_amount / total_transactions if total_transactions > 0 else 0
        
        currency = self.currency_combo.currentText()
        
        # İstatistik kartları
        stats = [
            ("Toplam İşlem Sayısı", f"{total_transactions:,}", "📊"),
            ("Toplam KDV Hariç Tutar", f"{self.format_currency(total_amount, currency)} {currency}", "💰"),
            ("Toplam KDV Tutarı", f"{self.format_currency(total_vat, currency)} {currency}", "💵"),
            ("En Yüksek İşlem", f"{self.format_currency(max_amount, currency)} {currency}", "📈"),
            ("Ortalama İşlem Tutarı", f"{self.format_currency(avg_amount, currency)} {currency}", "📉"),
        ]
        
        for idx, (label, value, icon) in enumerate(stats):
            card = QFrame()
            card.setStyleSheet("""
                QFrame {
                    background-color: #1E1E1E;
                    border-radius: 5px;
                    padding: 5px;
                }
            """)
            
            card_layout = QVBoxLayout()
            
            icon_label = QLabel(icon)
            icon_label.setStyleSheet("font-size: 13px;")
            icon_label.setAlignment(Qt.AlignCenter)
            
            label = QLabel(label)
            label.setStyleSheet("color: #888; font-size: 13px;")
            label.setAlignment(Qt.AlignCenter)
            
            value_label = QLabel(value)
            value_label.setStyleSheet("color: white; font-size: 15px; font-weight: bold;")
            value_label.setAlignment(Qt.AlignCenter)
            
            card_layout.addWidget(icon_label)
            card_layout.addWidget(label)
            card_layout.addWidget(value_label)
            card.setLayout(card_layout)
            
            stats_layout.addWidget(card, idx // 5, idx % 5)
        
        stats_group.setLayout(stats_layout)
        layout.addWidget(stats_group)
        
        # KDV Oranı Dağılımı Grafiği
        chart_group = QGroupBox("KDV Oranı Dağılımı")
        chart_group.setStyleSheet(stats_group.styleSheet())
        chart_layout = QVBoxLayout()
        
        figure = Figure(figsize=(8, 6), facecolor='#2D2D2D')
        canvas = FigureCanvas(figure)
        ax = figure.add_subplot(111)
        
        # KDV oranlarına göre işlem sayılarını hesapla
        kdv_rates = {}
        for row in range(total_transactions):
            rate = self.history_table.item(row, 3).text()
            kdv_rates[rate] = kdv_rates.get(rate, 0) + 1
        
        # Grafik çiz
        bars = ax.bar(kdv_rates.keys(), kdv_rates.values())
        ax.set_facecolor('#2D2D2D')
        ax.set_title('KDV Oranlarına Göre İşlem Dağılımı', color='white')
        ax.set_xlabel('KDV Oranı', color='white')
        ax.set_ylabel('İşlem Sayısı', color='white')
        ax.tick_params(colors='white')
        
        for spine in ax.spines.values():
            spine.set_color('#3D3D3D')
        
        # Barların üzerine değerleri yaz
        for bar in bars:
            height = bar.get_height()
            ax.text(bar.get_x() + bar.get_width()/2., height,
                    f'{int(height)}',
                    ha='center', va='bottom', color='white')
        
        chart_layout.addWidget(canvas)
        chart_group.setLayout(chart_layout)
        layout.addWidget(chart_group)
        
        # Dialog stil
        dialog.setStyleSheet("""
            QDialog {
                background-color: #1E1E1E;
            }
            QLabel {
                color: white;
            }
        """)
        
        dialog.setLayout(layout)
        dialog.exec_()

    def show_bulk_calc(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("Toplu Hesaplama")
        dialog.setModal(True)
        dialog.resize(800, 600)
        
        # Son seçilen değerleri yükle
        last_bulk_kdv = self.settings.value('last_bulk_kdv', '%18')
        last_bulk_currency = self.settings.value('last_bulk_currency', 'TRY')
        
        layout = QVBoxLayout()
        
        # Başlık
        title = QLabel("Toplu KDV Hesaplama")
        title.setStyleSheet("font-size: 18px; font-weight: bold; color: white;")
        title.setAlignment(Qt.AlignCenter)
        layout.addWidget(title)
        
        # KDV ve Para Birimi seçimi
        settings_group = QGroupBox("Hesaplama Ayarları")
        settings_group.setStyleSheet("color: white;")
        settings_layout = QHBoxLayout()
        
        kdv_label = QLabel("KDV Oranı:")
        kdv_combo = QComboBox()
        kdv_combo.addItems([self.kdv_combo.itemText(i) for i in range(self.kdv_combo.count())])
        # Son seçilen KDV oranını ayarla
        index = kdv_combo.findText(last_bulk_kdv)
        kdv_combo.setCurrentIndex(index if index >= 0 else 0)
        
        currency_label = QLabel("Para Birimi:")
        currency_combo = QComboBox()
        currency_combo.addItems(['TRY', 'USD', 'EUR', 'GBP'])
        # Son seçilen para birimini ayarla
        index = currency_combo.findText(last_bulk_currency)
        currency_combo.setCurrentIndex(index if index >= 0 else 0)
        
        settings_layout.addWidget(kdv_label)
        settings_layout.addWidget(kdv_combo)
        settings_layout.addWidget(currency_label)
        settings_layout.addWidget(currency_combo)
        settings_group.setLayout(settings_layout)
        layout.addWidget(settings_group)
        
        # Seçimleri kaydet
        def save_selections():
            self.settings.setValue('last_bulk_kdv', kdv_combo.currentText())
            self.settings.setValue('last_bulk_currency', currency_combo.currentText())
        
        kdv_combo.currentTextChanged.connect(save_selections)
        currency_combo.currentTextChanged.connect(save_selections)
        
        # Tutar girişi
        input_group = QGroupBox("Tutarlar")
        input_group.setStyleSheet("color: white;")
        input_layout = QVBoxLayout()
        
        text_edit = QTextEdit()
        text_edit.setPlaceholderText("Her satıra bir tutar giriniz...\nÖrnek:\n1000\n2500.50\n750.75")
        input_layout.addWidget(text_edit)
        input_group.setLayout(input_layout)
        layout.addWidget(input_group)
        
        # Sonuç tablosu
        table = QTableWidget()
        table.setColumnCount(5)
        table.setHorizontalHeaderLabels(["Tutar", "Para Birimi", "KDV Oranı", "KDV Tutarı", "Toplam"])
        layout.addWidget(table)
        
        # Toplamlar
        totals_group = QGroupBox("Toplamlar")
        totals_group.setStyleSheet("color: white;")
        totals_layout = QVBoxLayout()
        
        total_base_label = QLabel("Toplam KDV Hariç Tutar: 0.00")
        total_vat_label = QLabel("Toplam KDV Tutarı: 0.00")
        total_sum_label = QLabel("Genel Toplam KDV Dahil: 0.00")
        
        totals_layout.addWidget(total_base_label)
        totals_layout.addWidget(total_vat_label)
        totals_layout.addWidget(total_sum_label)
        totals_group.setLayout(totals_layout)
        layout.addWidget(totals_group)
        
        # Butonlar
        buttons = QHBoxLayout()
        
        calc_btn = QPushButton("Hesapla")
        export_btn = QPushButton("PDF'e Aktar")
        close_btn = QPushButton("Kapat")
        
        def calculate():
            try:
                table.setRowCount(0)
                lines = text_edit.toPlainText().strip().split('\n')
                kdv = float(kdv_combo.currentText().strip('%')) / 100
                curr = currency_combo.currentText()
                
                total_base = total_vat = total_sum = 0
                
                for line in lines:
                    if not line.strip():
                        continue
                        
                    try:
                        amount = float(line.strip())
                        vat = amount * kdv
                        sum_total = amount + vat
                        
                        total_base += amount
                        total_vat += vat
                        total_sum += sum_total
                        
                        row = table.rowCount()
                        table.insertRow(row)
                        
                        items = [
                            f"{self.format_currency(amount, curr)}",
                            curr,
                            f"%{kdv*100}",
                            f"{self.format_currency(vat, curr)}",
                            f"{self.format_currency(sum_total, curr)}"
                        ]
                        
                        for col, item in enumerate(items):
                            table.setItem(row, col, QTableWidgetItem(item))
                            
                    except ValueError:
                        continue
                
                # Toplamları güncelle
                total_base_label.setText(f"Toplam KDV Hariç Tutar: {self.format_currency(total_base, curr)} {curr}")
                total_vat_label.setText(f"Toplam KDV Tutarı: {self.format_currency(total_vat, curr)} {curr}")
                total_sum_label.setText(f"Genel Toplam KDV Dahil: {self.format_currency(total_sum, curr)} {curr}")
                
            except Exception as e:
                QMessageBox.critical(dialog, "Hata", str(e))
        
        def export_pdf():
            try:
                filename, _ = QFileDialog.getSaveFileName(dialog, "PDF'e Aktar", "", "PDF Dosyası (*.pdf)")
                if filename:
                    if not filename.endswith('.pdf'):
                        filename += '.pdf'
                        
                    printer = QPrinter(QPrinter.HighResolution)
                    printer.setOutputFormat(QPrinter.PdfFormat)
                    printer.setOutputFileName(filename)
                    
                    doc = QTextDocument()
                    
                    # Tablo ve toplamları HTML olarak oluştur
                    html = "<h2>Toplu KDV Hesaplama Raporu</h2><br>"
                    html += "<table border='1' cellspacing='0' cellpadding='3'>"
                    
                    # Tablo başlıkları
                    headers = ["Tutar", "Para Birimi", "KDV Oranı", "KDV Tutarı", "Toplam"]
                    html += "<tr>"
                    for header in headers:
                        html += f"<th>{header}</th>"
                    html += "</tr>"
                    
                    # Tablo verileri
                    for row in range(table.rowCount()):
                        html += "<tr>"
                        for col in range(table.columnCount()):
                            value = table.item(row, col).text() if table.item(row, col) else ''
                            html += f"<td align='center'>{value}</td>"
                        html += "</tr>"
                    
                    html += "</table><br><br>"
                    
                    # Toplamları ekle
                    html += f"""
                    <p><strong>{total_base_label.text()}</strong></p>
                    <p><strong>{total_vat_label.text()}</strong></p>
                    <p><strong>{total_sum_label.text()}</strong></p>
                    """
                    
                    doc.setHtml(html)
                    doc.print_(printer)
                    QMessageBox.information(dialog, "Başarılı", "PDF dosyası kaydedildi!")
                    
            except Exception as e:
                QMessageBox.critical(dialog, "Hata", f"PDF oluşturulurken hata: {str(e)}")
        
        calc_btn.clicked.connect(calculate)
        export_btn.clicked.connect(export_pdf)
        close_btn.clicked.connect(dialog.close)

        # Tablo özelliklerini ayarla
        table.setEditTriggers(QTableWidget.NoEditTriggers)

        # Seçili satırı silme butonu
        delete_row_btn = QPushButton("Seçili Satırı Sil")
        delete_row_btn.setStyleSheet("""
            QPushButton {
                background-color: #f44336;
                color: white;
                padding: 8px;
                border-radius: 4px;
            }
        """)

        def delete_selected_bulk_row():
            selected_rows = table.selectedIndexes()
            if selected_rows:
                selected_row = selected_rows[0].row()
                reply = QMessageBox.question(dialog, 'Onay',
                    'Seçili satır silinecek. Emin misiniz?',
                    QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
                
                if reply == QMessageBox.Yes:
                    table.removeRow(selected_row)
                    # Toplamları güncelle
                    try:
                        total_base = total_vat = total_sum = 0
                        curr = currency_combo.currentText()
                        
                        for row in range(table.rowCount()):
                            amount = float(table.item(row, 0).text().split()[0].replace(',', '.'))
                            vat = float(table.item(row, 3).text().split()[0].replace(',', '.'))
                            total = float(table.item(row, 4).text().split()[0].replace(',', '.'))
                            
                            total_base += amount
                            total_vat += vat
                            total_sum += total
                        
                        # Toplamları güncelle  
                        total_base_label.setText(f"Toplam KDV Hariç Tutar: {self.format_currency(total_base, curr)} {curr}")
                        total_vat_label.setText(f"Toplam KDV Tutarı: {self.format_currency(total_vat, curr)} {curr}")
                        total_sum_label.setText(f"Genel Toplam KDV Dahil: {self.format_currency(total_sum, curr)} {curr}")
                    except:
                        pass

        delete_row_btn.clicked.connect(delete_selected_bulk_row)

        # Layout güncelleme
        buttons.addWidget(delete_row_btn)
        buttons.addWidget(calc_btn)
        buttons.addWidget(export_btn)
        buttons.addWidget(close_btn)
        
        buttons.addWidget(calc_btn)
        buttons.addWidget(export_btn)
        buttons.addWidget(close_btn)
        
        buttons.addWidget(delete_row_btn)
        layout.addLayout(buttons)
        
        # Karanlık tema
        dialog.setStyleSheet("""
            QDialog { background-color: #1e1e1e; }
            QLabel { color: white; }
            QPushButton { 
                background-color: #0078d4;
                color: white;
                padding: 8px;
                border: none;
                border-radius: 4px;
            }
            QTableWidget {
                background-color: #2d2d2d;
                color: white;
                gridline-color: #3d3d3d;
            }
            QTableWidget QHeaderView::section {
                background-color: #2d2d2d;
                color: white;
                border: 1px solid #3d3d3d;
            }
            QComboBox, QTextEdit {
                background-color: #2d2d2d;
                color: white;
                border: 1px solid #3d3d3d;
                padding: 4px;
            }
        """)
        
        dialog.setLayout(layout)
        dialog.exec_()

    def change_theme(self, theme):
        # Tema değişikliğini kaydet
        self.settings.setValue('theme', theme)
        self.settings.sync()
        
        if theme == 'light':
            self.setStyleSheet(self.light_theme_style)
        else:
            self.setStyleSheet(self.dark_theme_style)

    def get_logo_path(self):
        """Logo dosyasının yolunu döndürür."""
        if platform.system() == "Linux":
            if hasattr(sys, "_MEIPASS"):
                return os.path.join(sys._MEIPASS, "kdv.png")
            elif os.path.exists("/usr/share/icons/hicolor/48x48/apps/kdv.png"):
                return "/usr/share/icons/hicolor/48x48/apps/kdv.png"
            elif os.path.exists("kdv.png"):
                return "kdv.png"
            return None
        else:
            return get_resource_path("kdv.png")

    def get_icon_path(self):
        """Simge dosyasının yolunu döndürür."""
        if platform.system() == "Linux":
            icon_paths = [
                os.path.join(os.path.dirname(__file__), "kdv.png"),
                "/usr/share/icons/hicolor/48x48/apps/kdv.png", 
                os.path.join(os.getcwd(), "kdv.png")
            ]
            
            for path in icon_paths:
                if os.path.exists(path):
                    return path
            return None
        else:
            return get_resource_path("kdv.png")

    def change_language(self, lang):
        self.settings.setValue('language', lang)
        self.settings.sync()
        self.current_language = lang
        self.update_texts()
        
    def update_texts(self):
        # Pencere başlığı
        self.setWindowTitle(self.tr('window_title'))
        
        # Menüleri güncelle
        self.menuBar().clear()
        self.create_menu()
        
        # Widget metinlerini güncelle
        self.tutar_label.setText(self.tr('amount'))
        self.tutar_input.setPlaceholderText(self.tr('amount_placeholder'))
        self.kdv_label.setText(self.tr('vat_rate'))
        self.hesapla_btn.setText(self.tr('calculate'))
        self.history_label.setText(self.tr('history'))
        self.clear_history_btn.setText(self.tr('clear_history'))
        
        # Diğer bileşenleri güncelle
        self.kdv_dahil_group.setTitle(self.tr('vat_included'))
        self.kdv_haric_group.setTitle(self.tr('vat_excluded'))

    def tr(self, key):
        return self.translations[self.current_language].get(key, key)

    def get_currency_decimals(self, currency):
        """Para birimine göre ondalık hane sayısını belirle"""
        currency_decimals = {
            'JPY': 0,  # Japon Yeni ondalık kullanmaz
            'TRY': 2,
            'USD': 2,
            'EUR': 2,
            'GBP': 2,
            'CHF': 2,
            'AUD': 2,
            'CAD': 2
        }
        return currency_decimals.get(currency, 2)

    def format_currency(self, value, currency):
        try:
            decimals = self.get_currency_decimals(currency)
            converted = self.exchange.convert(value, 'TRY', currency)
            return f"{converted:.{decimals}f}"
        except:
            return f"{value:.2f}"  # Hata durumunda basit formatlama

    def delete_selected_row(self):
        selected_rows = self.history_table.selectedItems()
        if selected_rows:
            reply = QMessageBox.question(self, 'Onay', 
                'Seçili satır silinecek. Emin misiniz?',
                QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
                
            if reply == QMessageBox.Yes:
                row = selected_rows[0].row()
                self.history_table.removeRow(row)
                self.update_totals()
                self.save_history_to_settings()

class KDVOranYonetimi(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.parent = parent
        self.setWindowTitle("KDV Oran Yönetimi")
        self.setModal(True)
        self.setMinimumWidth(400)
        self.setup_ui()
        
    def setup_ui(self):
        layout = QVBoxLayout()
        
        # Başlık etiketi
        title_label = QLabel("KDV Oran Yönetimi")
        title_label.setStyleSheet("""
            QLabel {
                font-size: 18px;
                font-weight: bold;
                color: #2196F3;
                padding: 10px;
            }
        """)
        title_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(title_label)
        
        # Mevcut oranlar grubu
        rates_group = QGroupBox("Mevcut KDV Oranları")
        rates_layout = QVBoxLayout()
        
        # Oran listesi
        self.rates_list = QListWidget()
        self.rates_list.setStyleSheet("""
            QListWidget {
                border: 1px solid #cccccc;
                border-radius: 4px;
                padding: 5px;
                background-color: #2D2D2D;
            }
            QListWidget::item {
                padding: 8px;
                border-bottom: 1px solid #eeeeee;
            }
            QListWidget::item:selected {
                background-color: #2196F3;
                color: white;
            }
        """)
        self.rates_list.addItems([self.parent.kdv_combo.itemText(i) 
                                for i in range(self.parent.kdv_combo.count())])
        rates_layout.addWidget(self.rates_list)
        rates_group.setLayout(rates_layout)
        layout.addWidget(rates_group)
        
        # Yeni oran ekleme grubu
        new_rate_group = QGroupBox("Yeni KDV Oranı Ekle")
        new_rate_layout = QHBoxLayout()
        
        self.new_rate_input = QLineEdit()
        self.new_rate_input.setPlaceholderText("Örn: 20")
        self.new_rate_input.setStyleSheet("""
            QLineEdit {
                padding: 8px;
                border: 1px solid #cccccc;
                border-radius: 4px;
                background-color: #2D2D2D;
            }
            QLineEdit:focus {
                border-color: #2196F3;
            }
        """)
        
        self.add_btn = QPushButton("Ekle")
        self.add_btn.setStyleSheet("""
            QPushButton {
                background-color: #2196F3;
                color: white;
                padding: 8px 16px;
                border-radius: 4px;
                border: none;
                font-weight: bold;
            }
            QPushButton:hover {
                background-color: #1976D2;
            }
            QPushButton:pressed {
                background-color: #0D47A1;
            }
        """)
        self.add_btn.clicked.connect(self.add_rate)
        
        new_rate_layout.addWidget(self.new_rate_input)
        new_rate_layout.addWidget(self.add_btn)
        new_rate_group.setLayout(new_rate_layout)
        layout.addWidget(new_rate_group)
        
        # İşlem butonları
        buttons_layout = QHBoxLayout()
        
        self.delete_btn = QPushButton("Seçili Oranı Sil")
        self.delete_btn.setStyleSheet("""
            QPushButton {
                background-color: #f44336;
                color: white;
                padding: 8px 16px;
                border-radius: 4px;
                border: none;
            }
            QPushButton:hover {
                background-color: #d32f2f;
            }
            QPushButton:pressed {
                background-color: #b71c1c;
            }
        """)
        self.delete_btn.clicked.connect(self.delete_rate)
        
        close_btn = QPushButton("Kapat")
        close_btn.setStyleSheet("""
            QPushButton {
                background-color: #757575;
                color: white;
                padding: 8px 16px;
                border-radius: 4px;
                border: none;
            }
            QPushButton:hover {
                background-color: #616161;
            }
            QPushButton:pressed {
                background-color: #424242;
            }
        """)
        close_btn.clicked.connect(self.close)
        
        buttons_layout.addWidget(self.delete_btn)
        buttons_layout.addWidget(close_btn)
        layout.addLayout(buttons_layout)
        
        self.setLayout(layout)
        
    def add_rate(self):
        try:
            rate = float(self.new_rate_input.text())
            if 0 < rate <= 100:
                rate_text = f"%{rate}"
                existing_rates = [self.rates_list.item(i).text() 
                                for i in range(self.rates_list.count())]
                
                if rate_text not in existing_rates:
                    self.rates_list.addItem(rate_text)
                    self.parent.kdv_combo.addItem(rate_text)
                    self.new_rate_input.clear()
                    self.parent.save_kdv_rates()
                    QMessageBox.information(self, "Başarılı", f"{rate_text} oranı başarıyla eklendi.")
                else:
                    QMessageBox.warning(self, "Hata", "Bu KDV oranı zaten mevcut!")
            else:
                QMessageBox.warning(self, "Hata", "KDV oranı 0-100 arasında olmalıdır!")
        except ValueError:
            QMessageBox.warning(self, "Hata", "Geçerli bir sayı giriniz!")
            
    def delete_rate(self):
        current_item = self.rates_list.currentItem()
        if current_item:
            reply = QMessageBox.question(self, "Onay", 
                                       f"{current_item.text()} oranını silmek istediğinizden emin misiniz?",
                                       QMessageBox.Yes | QMessageBox.No)
            
            if reply == QMessageBox.Yes:
                rate_text = current_item.text()
                # Listeden kaldır
                self.rates_list.takeItem(self.rates_list.row(current_item))
                # Combo box'tan kaldır
                index = self.parent.kdv_combo.findText(rate_text)
                if index >= 0:
                    self.parent.kdv_combo.removeItem(index)
                self.parent.save_kdv_rates()
                QMessageBox.information(self, "Başarılı", f"{rate_text} oranı başarıyla silindi.")
        else:
            QMessageBox.warning(self, "Hata", "Lütfen silmek istediğiniz oranı seçin!")

if __name__ == '__main__':
    app = QApplication(sys.argv)
    
    # Platform özel ayarlar
    if platform.system() == "Darwin":  # macOS
        app.setAttribute(Qt.AA_DontShowIconsInMenus, False)
    
    # Uygulama ikonu
    icon_path = get_resource_path("kdv.png")
    if icon_path:
        app.setWindowIcon(QIcon(icon_path))
    
    window = KDVHesaplama()
    window.show()
    sys.exit(app.exec_())